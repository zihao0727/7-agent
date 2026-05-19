from __future__ import annotations

import csv
import html
import io
import re
import uuid
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import httpx

from backend.db import get_db
from backend.session_access import session_filter_for_user

MAX_FILE_BYTES = 40 * 1024 * 1024
CHUNK_SIZE = 1400
CHUNK_OVERLAP = 180
ACTIVE_STATUS = "active"


def utc_now() -> datetime:
    return datetime.utcnow()


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def tokenize(text: str) -> list[str]:
    text = text.lower()
    latin = re.findall(r"[a-z0-9_]{2,}", text)
    cjk_chunks = re.findall(r"[\u4e00-\u9fff]{2,}", text)
    cjk_pairs: list[str] = []
    for chunk in cjk_chunks:
        cjk_pairs.extend(chunk[i : i + 2] for i in range(max(0, len(chunk) - 1)))
    return sorted(set(latin + cjk_pairs))[:300]


def serialize_doc(doc: dict[str, Any]) -> dict[str, Any]:
    row = dict(doc)
    row["id"] = str(row.get("_id") or row.get("id"))
    row.pop("_id", None)
    for key in ("created_at", "updated_at"):
        if isinstance(row.get(key), datetime):
            row[key] = row[key].isoformat()
    return row


class TextOnlyHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.in_ignored = False
        self.title = ""
        self._in_title = False
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "noscript"}:
            self.in_ignored = True
        if tag == "title":
            self._in_title = True
        if tag in {"p", "br", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "noscript"}:
            self.in_ignored = False
        if tag == "title":
            self._in_title = False
        if tag in {"p", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self.in_ignored:
            return
        text = html.unescape(data).strip()
        if not text:
            return
        if self._in_title:
            self.title += text
        self.parts.append(text)

    def get_text(self) -> str:
        return normalize_text("\n".join(self.parts))


def read_text_bytes(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def extract_pdf(raw: bytes) -> list[dict[str, Any]]:
    import pypdf  # type: ignore

    reader = pypdf.PdfReader(io.BytesIO(raw))
    sections: list[dict[str, Any]] = []
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            sections.append({"text": text.strip(), "page": idx + 1})
    return sections


def extract_docx(raw: bytes) -> list[dict[str, Any]]:
    import docx  # type: ignore

    doc = docx.Document(io.BytesIO(raw))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [{"text": "\n".join(parts)}]


def extract_xlsx(raw: bytes) -> list[dict[str, Any]]:
    import openpyxl  # type: ignore

    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sections: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows: list[str] = []
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx >= 300:
                rows.append("... truncated after 300 rows ...")
                break
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append("\t".join(values))
        if rows:
            sections.append({"text": "\n".join(rows), "section": sheet.title})
    return sections


def extract_csv_text(raw: bytes, suffix: str) -> list[dict[str, Any]]:
    text = read_text_bytes(raw)
    delimiter = "\t" if suffix == ".tsv" else ","
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = []
        for idx, row in enumerate(reader):
            if idx >= 600:
                rows.append("... truncated after 600 rows ...")
                break
            rows.append("\t".join(row))
        return [{"text": "\n".join(rows)}]
    except Exception:
        return [{"text": text}]


def extract_file_sections(filename: str, content_type: str, raw: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("File is too large for knowledge ingestion")
    if suffix == ".pdf" or content_type == "application/pdf":
        return extract_pdf(raw)
    if suffix == ".docx":
        return extract_docx(raw)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(raw)
    if suffix in {".csv", ".tsv"}:
        return extract_csv_text(raw, suffix)
    if suffix in {".html", ".htm"} or "html" in content_type:
        parser = TextOnlyHTMLParser()
        parser.feed(read_text_bytes(raw))
        return [{"text": parser.get_text(), "section": parser.title.strip() or None}]
    return [{"text": read_text_bytes(raw)}]


def split_sections(sections: list[dict[str, Any]], source_name: str) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for section in sections:
        text = normalize_text(str(section.get("text") or ""))
        if not text:
            continue
        start = 0
        while start < len(text):
            part = text[start : start + CHUNK_SIZE].strip()
            if part:
                source_bits = [source_name]
                if section.get("page"):
                    source_bits.append(f"p.{section['page']}")
                if section.get("section"):
                    source_bits.append(str(section["section"]))
                chunks.append(
                    {
                        "text": part,
                        "page": section.get("page"),
                        "section": section.get("section"),
                        "source_label": " / ".join(source_bits),
                    }
                )
            if start + CHUNK_SIZE >= len(text):
                break
            start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks


async def create_document(
    *,
    user_id: int,
    title: str,
    source_type: str,
    sections: list[dict[str, Any]],
    filename: str | None = None,
    content_type: str | None = None,
    url: str | None = None,
    session_id: str | None = None,
) -> dict[str, Any]:
    db = get_db()
    doc_id = str(uuid.uuid4())
    now = utc_now()
    source_name = filename or title or url or doc_id
    chunks = split_sections(sections, source_name)
    doc = {
        "_id": doc_id,
        "user_id": int(user_id),
        "title": title or filename or url or "Untitled document",
        "source_type": source_type,
        "filename": filename,
        "content_type": content_type,
        "url": url,
        "session_id": session_id,
        "status": ACTIVE_STATUS,
        "chunk_count": len(chunks),
        "created_at": now,
        "updated_at": now,
    }
    await db["knowledge_documents"].insert_one(doc)
    if chunks:
        await db["knowledge_chunks"].insert_many(
            [
                {
                    "_id": str(uuid.uuid4()),
                    "user_id": int(user_id),
                    "document_id": doc_id,
                    "chunk_index": idx,
                    "text": chunk["text"],
                    "tokens": tokenize(chunk["text"]),
                    "source_label": chunk["source_label"],
                    "page": chunk.get("page"),
                    "section": chunk.get("section"),
                    "created_at": now,
                }
                for idx, chunk in enumerate(chunks)
            ]
        )
    return serialize_doc(doc)


async def ingest_file(
    *,
    user_id: int,
    filename: str,
    content_type: str,
    raw: bytes,
) -> dict[str, Any]:
    sections = extract_file_sections(filename, content_type, raw)
    return await create_document(
        user_id=user_id,
        title=filename,
        source_type="file",
        filename=filename,
        content_type=content_type,
        sections=sections,
    )


async def ingest_url(*, user_id: int, url: str) -> dict[str, Any]:
    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as client:
        resp = await client.get(url)
        resp.raise_for_status()
    parser = TextOnlyHTMLParser()
    parser.feed(resp.text)
    title = normalize_text(parser.title) or url
    return await create_document(
        user_id=user_id,
        title=title[:200],
        source_type="url",
        url=url,
        content_type=resp.headers.get("content-type"),
        sections=[{"text": parser.get_text(), "section": title}],
    )


async def ingest_session(*, user_id: int, session_id: str) -> dict[str, Any]:
    db = get_db()
    session = await db["sessions"].find_one(session_filter_for_user(session_id, user_id))
    if not session:
        raise ValueError("Session not found")
    lines: list[str] = []
    for message in session.get("messages", []):
        role = message.get("role") or "unknown"
        content = normalize_text(str(message.get("content") or ""))
        if content:
            lines.append(f"{role}: {content}")
    return await create_document(
        user_id=user_id,
        title=f"Chat: {session.get('title') or session_id}",
        source_type="chat",
        session_id=session_id,
        sections=[{"text": "\n".join(lines)}],
    )


async def list_documents(user_id: int) -> list[dict[str, Any]]:
    rows = (
        await get_db()["knowledge_documents"]
        .find({"user_id": int(user_id), "status": ACTIVE_STATUS})
        .sort("updated_at", -1)
        .to_list(None)
    )
    return [serialize_doc(row) for row in rows]


async def delete_document(user_id: int, document_id: str) -> bool:
    db = get_db()
    result = await db["knowledge_documents"].update_one(
        {"_id": document_id, "user_id": int(user_id)},
        {"$set": {"status": "deleted", "updated_at": utc_now()}},
    )
    if result.modified_count:
        await db["knowledge_chunks"].delete_many(
            {"document_id": document_id, "user_id": int(user_id)}
        )
    return result.modified_count > 0


async def search_knowledge(user_id: int, query: str, limit: int = 6) -> list[dict[str, Any]]:
    db = get_db()
    q = normalize_text(query)
    q_tokens = set(tokenize(q))
    if q_tokens:
        rows = (
            await db["knowledge_chunks"]
            .find({"user_id": int(user_id), "tokens": {"$in": list(q_tokens)}})
            .limit(250)
            .to_list(None)
        )
    else:
        rows = (
            await db["knowledge_chunks"]
            .find({"user_id": int(user_id)})
            .sort("created_at", -1)
            .limit(80)
            .to_list(None)
        )
    if not rows:
        return []

    doc_ids = list({row["document_id"] for row in rows})
    docs = {
        str(doc["_id"]): doc
        for doc in await db["knowledge_documents"]
        .find({"_id": {"$in": doc_ids}, "user_id": int(user_id), "status": ACTIVE_STATUS})
        .to_list(None)
    }

    scored: list[tuple[float, dict[str, Any]]] = []
    q_lower = q.lower()
    for row in rows:
        doc = docs.get(row["document_id"])
        if not doc:
            continue
        tokens = set(row.get("tokens") or [])
        overlap = len(q_tokens & tokens)
        text = str(row.get("text") or "")
        score = overlap * 4.0
        if q_lower and q_lower in text.lower():
            score += 8.0
        score += min(len(text), 1400) / 1400
        scored.append((score, {**row, "document": doc}))

    scored.sort(key=lambda item: item[0], reverse=True)
    results: list[dict[str, Any]] = []
    for score, row in scored[: max(1, min(limit, 20))]:
        doc = row["document"]
        results.append(
            {
                "document_id": row["document_id"],
                "document_title": doc.get("title"),
                "source_type": doc.get("source_type"),
                "url": doc.get("url"),
                "chunk_index": row.get("chunk_index"),
                "source_label": row.get("source_label"),
                "score": round(score, 3),
                "text": row.get("text"),
            }
        )
    return results


def format_search_results(results: list[dict[str, Any]]) -> str:
    if not results:
        return "No knowledge-base matches were found."
    lines = ["Knowledge-base matches:"]
    for idx, item in enumerate(results, 1):
        citation = f"{item.get('source_label') or item.get('document_title')} #chunk-{item.get('chunk_index')}"
        if item.get("url"):
            citation += f" ({item['url']})"
        lines.append(f"\n[{idx}] {citation}")
        lines.append(str(item.get("text") or "")[:1600])
    lines.append("\nWhen answering, cite the bracketed match numbers and source labels.")
    return "\n".join(lines)
